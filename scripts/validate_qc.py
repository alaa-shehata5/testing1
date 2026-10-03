#!/usr/bin/env python3
"""Run repeatable Phase 14 consistency checks for the QA portfolio."""

from __future__ import annotations

import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import unquote

try:
    from openpyxl import load_workbook
except ImportError as exc:
    raise SystemExit(
        "Missing QC dependency. Run: "
        "python3 -m pip install -r scripts/requirements-qc.txt"
    ) from exc


ROOT = Path(__file__).resolve().parents[1]
EXPECTED_STATUSES = Counter(
    {"Pass": 17, "Fail": 2, "Blocked": 2, "Not Run": 33}
)
EXPECTED_TYPES = Counter(
    {
        "Positive": 20,
        "Negative": 14,
        "Boundary": 6,
        "State": 5,
        "UI": 4,
        "Compatibility": 3,
        "Accessibility": 2,
    }
)
errors: list[str] = []


def check(condition: bool, message: str) -> None:
    if not condition:
        errors.append(message)


def worksheet(relative_path: str, sheet_name: str):
    workbook = load_workbook(ROOT / relative_path, read_only=True, data_only=True)
    return workbook, workbook[sheet_name]


def headers(sheet) -> dict[str, int]:
    return {cell.value: index for index, cell in enumerate(sheet[1])}


def rows_by_id(sheet, id_column: str) -> tuple[dict[str, int], dict[str, tuple]]:
    column_map = headers(sheet)
    records = {
        row[column_map[id_column]]: row
        for row in sheet.iter_rows(min_row=2, values_only=True)
    }
    return column_map, records


def requirement_case_ids(value: str) -> list[str]:
    result: list[str] = []
    pattern = re.compile(r"(TC-[A-Z]+-)(\d{3})(?:[–-](\d{3}))?")
    for prefix, first, last in pattern.findall(str(value)):
        start = int(first)
        end = int(last) if last else start
        result.extend(f"{prefix}{number:03d}" for number in range(start, end + 1))
    return result


def check_markdown_links(tracked_files: list[str]) -> None:
    link_pattern = re.compile(r"!?\[[^\]]*\]\(([^)]+)\)")
    for relative in tracked_files:
        if Path(relative).suffix.lower() != ".md":
            continue
        content = (ROOT / relative).read_text(encoding="utf-8")
        for raw_target in link_pattern.findall(content):
            target = raw_target.strip().split()[0].strip("<>")
            if target.startswith(("https://", "http://", "mailto:", "#")):
                continue
            path_part = unquote(target.split("#", 1)[0].split("?", 1)[0])
            if not path_part:
                continue
            target_path = (ROOT / relative).parent / path_part
            check(
                target_path.exists(),
                f"Broken local Markdown link: {relative} -> {target}",
            )


tracked_files = subprocess.check_output(
    ["git", "ls-files", "-z"], cwd=ROOT
).decode().split("\0")
tracked_files = [path for path in tracked_files if path]
tracked_set = set(tracked_files)

# Requirements and assumption separation.
drs = (ROOT / "docs/02-requirements/derived-requirements.md").read_text(
    encoding="utf-8"
)
requirement_ids = re.findall(r"^\| (REQ-[A-Z]+-\d{3}) \|", drs, re.MULTILINE)
assumptions = (ROOT / "docs/02-requirements/assumptions.md").read_text(
    encoding="utf-8"
)
assumption_ids = re.findall(r"ASM-[A-Z]+-\d{3}", assumptions)
check(len(requirement_ids) == 20 and len(set(requirement_ids)) == 20,
      "DRS must contain 20 unique requirement rows.")
check(len(assumption_ids) == 4 and len(set(assumption_ids)) == 4,
      "Assumptions file must contain the four unique ASM IDs.")
check(not re.search(r"^\| ASM-[A-Z]+-\d{3} \|", drs, re.MULTILINE),
      "An ASM row is present in the DRS table.")

# Test design, case execution, and status reconciliation.
case_book, cases_sheet = worksheet("docs/04-test-cases/test-cases.xlsx", "TestCases")
execution_book, execution_sheet = worksheet(
    "docs/06-execution/execution-results.xlsx", "Execution"
)
case_columns, cases = rows_by_id(cases_sheet, "TC ID")
execution_columns, executions = rows_by_id(execution_sheet, "TC ID")
check(len(cases) == 54, "Test-case workbook must contain 54 unique cases.")
check(set(cases) == set(executions),
      "Case and execution workbooks must contain identical TC IDs.")
check(
    Counter(row[case_columns["Test Type"]] for row in cases.values())
    == EXPECTED_TYPES,
    "Test-case type mix does not match the approved 54-case allocation.",
)
for tc_id, row in cases.items():
    for column in ("Requirement ID", "Test Data", "Steps", "Expected Result"):
        check(bool(row[case_columns[column]]), f"{tc_id} has an empty {column}.")
    if tc_id in executions:
        check(
            row[case_columns["Status"]]
            == executions[tc_id][execution_columns["Status"]],
            f"{tc_id} status differs between case and execution workbooks.",
        )
status_counts = Counter(
    row[execution_columns["Status"]] for row in executions.values()
)
check(status_counts == EXPECTED_STATUSES,
      f"Unexpected execution status totals: {dict(status_counts)}")
check(
    all(not row[execution_columns["Defect ID"]] for row in executions.values()),
    "Execution workbook contains a defect link despite the reported zero.",
)

# Requirement-to-case-to-execution traceability.
rtm_book, rtm_sheet = worksheet("docs/05-traceability/rtm.xlsx", "RTM")
rtm_columns, rtm_rows = rows_by_id(rtm_sheet, "Requirement ID")
expected_by_requirement: dict[str, set[str]] = defaultdict(set)
for tc_id, row in cases.items():
    expected_by_requirement[row[case_columns["Requirement ID"]]].add(tc_id)
check(set(rtm_rows) == set(requirement_ids),
      "RTM requirement IDs do not match the DRS.")
all_rtm_cases: list[str] = []
for req_id, row in rtm_rows.items():
    mapped = set(requirement_case_ids(row[rtm_columns["Test Case IDs"]]))
    all_rtm_cases.extend(mapped)
    check(mapped == expected_by_requirement[req_id],
          f"{req_id} RTM case mapping differs from the test-case workbook.")
    expected_counts = Counter(
        executions[tc_id][execution_columns["Status"]]
        for tc_id in mapped
        if tc_id in executions
    )
    recorded_counts = "; ".join(
        f"{expected_counts[status]} {status}"
        for status in ("Pass", "Fail", "Blocked", "N/A", "Not Run")
        if expected_counts[status]
    )
    check(
        row[rtm_columns["Execution"]] == recorded_counts,
        f"{req_id} RTM execution counts are stale.",
    )
check(Counter(all_rtm_cases) == Counter(cases.keys()),
      "Each test case must appear exactly once in the RTM.")
check(
    all(not row[rtm_columns["Defect ID"]] for row in rtm_rows.values()),
    "RTM contains a defect link despite the reported zero.",
)

bug_reports = [
    path for path in tracked_files
    if path.startswith("bugs/") and path.endswith("/bug-report.md")
]
check(len(bug_reports) == 1,
      f"Expected one explicitly unfilled bug template, found {len(bug_reports)}.")
if bug_reports:
    template = (ROOT / bug_reports[0]).read_text(encoding="utf-8").lower()
    check("unfilled template" in template,
          "The only bug report is not clearly marked as an unfilled template.")

# Expected regression evidence must resolve to repository files.
regression_book, regression_sheet = worksheet(
    "docs/06-execution/regression-results.xlsx", "Regression"
)
regression_columns = headers(regression_sheet)
for row in regression_sheet.iter_rows(min_row=2, values_only=True):
    evidence = row[regression_columns["Evidence"]]
    if not evidence:
        continue
    for evidence_path in str(evidence).split("; "):
        check(
            (ROOT / evidence_path).is_file(),
            f"Regression evidence path is missing: {evidence_path}",
        )

# Inventory: only the explicitly required PNG/XLSX artifacts are permitted
# binary deliverables; this is not a substitute for a secrets scanner.
pngs = [path for path in tracked_files if path.lower().endswith(".png")]
xlsx = [path for path in tracked_files if path.lower().endswith(".xlsx")]
other_binary_suffixes = {
    ".pdf", ".jpg", ".jpeg", ".gif", ".mp4", ".mov", ".zip", ".exe", ".bin",
    ".webp", ".7z", ".tar", ".gz", ".doc", ".docx", ".ppt", ".pptx",
}
unexpected_binaries = [
    path for path in tracked_files
    if Path(path).suffix.lower() in other_binary_suffixes
]
for path in tracked_files:
    if path in pngs or path in xlsx:
        continue
    content = (ROOT / path).read_bytes()
    try:
        content.decode("utf-8")
        looks_binary = b"\0" in content
    except UnicodeDecodeError:
        looks_binary = True
    if looks_binary and path not in unexpected_binaries:
        unexpected_binaries.append(path)
check(len(pngs) == 43, f"Expected 43 evidence PNGs, found {len(pngs)}.")
check(len(xlsx) == 5, f"Expected 5 QA workbooks, found {len(xlsx)}.")
check(all(path.startswith("evidence/") for path in pngs),
      "A PNG is tracked outside the evidence directory.")
check(all(path.startswith("docs/") for path in xlsx),
      "An XLSX workbook is tracked outside docs/.")
check(not unexpected_binaries,
      f"Unexpected tracked binary artifacts: {unexpected_binaries}")
check(".env" not in tracked_set, "Credential file .env must not be tracked.")
ignore_result = subprocess.run(
    ["git", "check-ignore", "-q", ".env"], cwd=ROOT, check=False
)
check(ignore_result.returncode == 0, ".env must remain ignored by Git.")
check(sum(path.startswith("test-data/") for path in tracked_files) == 8,
      "Expected eight tracked test-data files.")

# All repository-local links are checked, rather than sampling a subset.
check_markdown_links(tracked_files)

# Public-facing ethics and current measured totals.
for relative in ("README.md", "docs/07-final-report/final-qa-report.md",
                 "demo/case-study.md"):
    content = (ROOT / relative).read_text(encoding="utf-8").lower()
    check(
        "not commissioned by orangehrm" in content,
        f"Required portfolio disclaimer missing from {relative}.",
    )
expected_metric_rows = {
    "README.md": (
        "54 planned — 21 executed (38.9%): 17 Pass, 2 Fail*, "
        "2 Blocked, 33 Not Run"
    ),
    "docs/06-execution/test-summary.md": (
        "| Passed | 17 |",
        "| Failed | 2 ",
        "| Blocked | 2 ",
        "| Not Run | 33 ",
    ),
    "docs/07-final-report/final-qa-report.md": (
        "| Total planned | 54 |",
        "| Executed (Pass+Fail+Blocked) | 21 | 38.9% |",
        "| Passed | 17 | 81.0% of executed |",
        "| Failed | 2 |",
        "| Blocked | 2 |",
        "| Not Run | 33 |",
    ),
}
for relative, expected_fragments in expected_metric_rows.items():
    content = (ROOT / relative).read_text(encoding="utf-8")
    if isinstance(expected_fragments, str):
        expected_fragments = (expected_fragments,)
    for fragment in expected_fragments:
        check(fragment in content,
              f"Expected execution metric missing from {relative}: {fragment}")

case_book.close()
execution_book.close()
rtm_book.close()
regression_book.close()

if errors:
    print("QC validation FAILED:")
    for message in errors:
        print(f"- {message}")
    raise SystemExit(1)

print(
    "QC validation PASSED: "
    f"{len(requirement_ids)} requirements, {len(cases)} test cases, "
    f"{dict(status_counts)}, {len(pngs)} evidence PNGs, "
    f"{len(xlsx)} QA workbooks, {len(tracked_files)} tracked files; "
    "all repository-local Markdown links and regression evidence paths resolve."
)
